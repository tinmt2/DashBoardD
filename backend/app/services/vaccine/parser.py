"""
Vaccine spreadsheet parser.

Vaccine exports are named ``*.xls`` but are actually *Microsoft Spreadsheet
XML 2003* (``<ss:Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet">``).
The format is therefore detected from the file **content** (magic bytes /
XML root), never from the extension:

    detect file format -> format-specific reader -> worksheets -> header -> rows

Supported formats
-----------------
* ``SPREADSHEET_XML`` – native stdlib XML reader (no extra dependency).
* ``XLSX``            – via openpyxl (already a backend dependency).
* ``XLS`` (BIFF/OLE2) – only if ``xlrd`` is installed; otherwise a clear error.
* HTML-table "xls" exports and unknown content are rejected with a clear error.

The parser is purely structural: it returns typed cell values
(``str`` / ``float`` / ``int`` / ``bool`` / ``datetime`` / ``None``) and does
not interpret business meaning. Mapping is done by ``VaccineColumnMapper``.
"""
from __future__ import annotations

import io
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, List, Optional

SS_NS = "urn:schemas-microsoft-com:office:spreadsheet"
_NS = "{" + SS_NS + "}"

FORMAT_SPREADSHEET_XML = "SPREADSHEET_XML"
FORMAT_XLSX = "XLSX"
FORMAT_XLS = "XLS"
FORMAT_HTML = "HTML"
FORMAT_UNKNOWN = "UNKNOWN"

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"


class VaccineParseError(Exception):
    """File-level failure: the file cannot be read as a Vaccine report."""


@dataclass
class ParsedRow:
    row_number: int  # 1-based row number as seen in Excel
    values: List[Any]


@dataclass
class ParsedSheet:
    name: str
    header_row_number: int = 0
    headers: List[str] = field(default_factory=list)
    rows: List[ParsedRow] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    file_format: str
    sheets: List[ParsedSheet]


class VaccineSpreadsheetParser:
    # ------------------------------------------------------------------ detect
    @staticmethod
    def detect_format(content: bytes) -> str:
        if not content:
            return FORMAT_UNKNOWN
        if content.startswith(_ZIP_MAGIC):
            return FORMAT_XLSX
        if content.startswith(_OLE2_MAGIC):
            return FORMAT_XLS

        head = content[:4096]
        # Strip UTF-8 BOM / leading whitespace before sniffing XML/HTML.
        if head.startswith(b"\xef\xbb\xbf"):
            head = head[3:]
        sniff = head.lstrip().lower()
        if sniff.startswith(b"<"):
            if SS_NS.encode() in head.lower() or b"<ss:workbook" in sniff or b"<workbook" in sniff:
                return FORMAT_SPREADSHEET_XML
            if b"<html" in sniff or b"<table" in sniff or b"<!doctype html" in sniff:
                return FORMAT_HTML
        return FORMAT_UNKNOWN

    # ------------------------------------------------------------------- parse
    def parse(self, content: bytes) -> ParsedWorkbook:
        fmt = self.detect_format(content)
        if fmt == FORMAT_SPREADSHEET_XML:
            sheets = self._parse_spreadsheet_xml(content)
        elif fmt == FORMAT_XLSX:
            sheets = self._parse_xlsx(content)
        elif fmt == FORMAT_XLS:
            sheets = self._parse_legacy_xls(content)
        elif fmt == FORMAT_HTML:
            raise VaccineParseError(
                "File là bảng HTML (không phải Excel). Vui lòng xuất lại báo cáo Vaccine dạng Excel."
            )
        else:
            raise VaccineParseError(
                "Không nhận diện được định dạng file. Hỗ trợ: Spreadsheet XML (.xls xuất từ Vaccine), .xlsx."
            )
        return ParsedWorkbook(file_format=fmt, sheets=sheets)

    # ------------------------------------------------------ Spreadsheet XML 2003
    def _parse_spreadsheet_xml(self, content: bytes) -> List[ParsedSheet]:
        # Defensive: Spreadsheet XML never needs a DTD. Rejecting DOCTYPE blocks
        # entity-expansion / external-entity tricks without extra dependencies.
        if b"<!DOCTYPE" in content[:8192].upper():
            raise VaccineParseError("File Spreadsheet XML chứa DOCTYPE không được hỗ trợ.")
        import re
        content = re.sub(b'&(?!(?:[a-zA-Z][a-zA-Z0-9]+|#[0-9]+|#x[0-9a-fA-F]+);)', b'&amp;', content)
        try:
            root = ET.fromstring(content)
        except ET.ParseError as exc:
            raise VaccineParseError(f"File Spreadsheet XML bị lỗi cấu trúc: {exc}") from exc

        if root.tag != f"{_NS}Workbook":
            raise VaccineParseError("Không tìm thấy <Workbook> trong file Spreadsheet XML.")

        sheets: List[ParsedSheet] = []
        for ws in root.findall(f"{_NS}Worksheet"):
            name = ws.get(f"{_NS}Name") or f"Sheet{len(sheets) + 1}"
            table = ws.find(f"{_NS}Table")
            grid = self._read_xml_table(table) if table is not None else []
            sheets.append(self._build_sheet(name, grid))
        if not sheets:
            raise VaccineParseError("File không có worksheet nào.")
        return sheets

    def _read_xml_table(self, table: ET.Element) -> List[ParsedRow]:
        """Read <Row>/<Cell> honouring ss:Index (sparse rows/cells) and ss:MergeAcross.

        Stray text between rows (the real Vaccine export contains a literal
        ``undefined`` after the header row) is ignored because only <Row>
        elements are visited.
        """
        rows: List[ParsedRow] = []
        current_row = 0
        for row_el in table.findall(f"{_NS}Row"):
            idx = row_el.get(f"{_NS}Index")
            current_row = int(idx) if idx and idx.isdigit() else current_row + 1

            values: List[Any] = []
            col = 0  # 0-based next column
            for cell in row_el.findall(f"{_NS}Cell"):
                cidx = cell.get(f"{_NS}Index")
                if cidx and cidx.isdigit():
                    col = int(cidx) - 1
                while len(values) < col:
                    values.append(None)
                values.append(self._read_xml_cell(cell))
                col += 1
                merge = cell.get(f"{_NS}MergeAcross")
                if merge and merge.isdigit():
                    for _ in range(int(merge)):
                        values.append(None)
                        col += 1
            rows.append(ParsedRow(row_number=current_row, values=values))
        return rows

    @staticmethod
    def _read_xml_cell(cell: ET.Element) -> Any:
        data = cell.find(f"{_NS}Data")
        if data is None:
            return None
        # itertext() also captures rich text (<html:Font>/<html:B> children).
        text = "".join(data.itertext())
        dtype = (data.get(f"{_NS}Type") or "String").strip()
        if dtype == "Number":
            try:
                num = float(text)
                return int(num) if num.is_integer() and "e" not in text.lower() and "." not in text else num
            except ValueError:
                return text
        if dtype == "DateTime":
            parsed = _parse_iso_datetime(text)
            return parsed if parsed is not None else text
        if dtype == "Boolean":
            return text.strip() in ("1", "true", "TRUE")
        return text  # String / Error / unknown -> keep raw text

    # -------------------------------------------------------------------- XLSX
    def _parse_xlsx(self, content: bytes) -> List[ParsedSheet]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - openpyxl is in requirements
            raise VaccineParseError("Thiếu thư viện openpyxl để đọc .xlsx") from exc
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise VaccineParseError(f"Không đọc được file .xlsx: {exc}") from exc
        sheets: List[ParsedSheet] = []
        try:
            for ws in wb.worksheets:
                grid = [
                    ParsedRow(row_number=i, values=list(r))
                    for i, r in enumerate(ws.iter_rows(values_only=True), start=1)
                ]
                sheets.append(self._build_sheet(ws.title, grid))
        finally:
            wb.close()
        return sheets

    # ------------------------------------------------------------- legacy XLS
    def _parse_legacy_xls(self, content: bytes) -> List[ParsedSheet]:
        try:
            import xlrd  # type: ignore
        except ImportError as exc:
            raise VaccineParseError(
                "File là Excel 97-2003 (BIFF) nhưng server chưa cài 'xlrd'. "
                "Vui lòng lưu lại file dạng .xlsx hoặc cài xlrd."
            ) from exc
        try:
            book = xlrd.open_workbook(file_contents=content)
        except Exception as exc:
            raise VaccineParseError(f"Không đọc được file .xls: {exc}") from exc
        sheets: List[ParsedSheet] = []
        for sh in book.sheets():
            grid: List[ParsedRow] = []
            for r in range(sh.nrows):
                vals: List[Any] = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        vals.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        vals.append(None)
                    else:
                        vals.append(cell.value)
                grid.append(ParsedRow(row_number=r + 1, values=vals))
            sheets.append(self._build_sheet(sh.name, grid))
        return sheets

    # ----------------------------------------------------------------- helpers
    @staticmethod
    def _is_blank(v: Any) -> bool:
        return v is None or (isinstance(v, str) and v.strip() == "")

    def _build_sheet(self, name: str, grid: List[ParsedRow]) -> ParsedSheet:
        """Header = first non-empty row; data = following non-empty rows."""
        sheet = ParsedSheet(name=name)
        header_idx: Optional[int] = None
        for i, row in enumerate(grid):
            if any(not self._is_blank(v) for v in row.values):
                header_idx = i
                break
        if header_idx is None:
            return sheet

        header_row = grid[header_idx]
        headers = ["" if v is None else str(v).strip() for v in header_row.values]
        while headers and headers[-1] == "":
            headers.pop()
        sheet.header_row_number = header_row.row_number
        sheet.headers = headers

        width = len(headers)
        for row in grid[header_idx + 1:]:
            if all(self._is_blank(v) for v in row.values):
                continue
            vals = list(row.values[:width]) + [None] * max(0, width - len(row.values))
            sheet.rows.append(ParsedRow(row_number=row.row_number, values=vals))
        return sheet


def _parse_iso_datetime(text: str) -> Optional[datetime]:
    """Spreadsheet XML DateTime cells look like ``2026-10-06T00:00:00.000``."""
    t = (text or "").strip()
    if not t:
        return None
    for fmt in ("%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(t, fmt)
        except ValueError:
            continue
    return None
